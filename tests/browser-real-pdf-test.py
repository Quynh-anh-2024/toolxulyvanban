"""Test TRÌNH DUYỆT THẬT với PDF thật từ hệ thống trường học (tests/real/):

Bài học 21/10/2026: pdf.js trên trình duyệt thiếu CMap/font chuẩn sẽ giải mã
sai font CID của PDF in từ cổng trường học — lớp chữ hoàn hảo bị tưởng là
"lỗi font" và app OCR oan, cho ra chữ sai dấu ("ý kiến" -> "ý kiện").
Test Node KHÔNG tái hiện được lỗi này (Node có sẵn CMap cạnh gói cài đặt),
nên mọi khẳng định về PDF hệ thống phải qua test trình duyệt này.

Từ 23/10/2026 giao diện TỐI GIẢN: không còn bản xem trước — mọi kiểm chứng
đi qua FILE TẢI XUỐNG thật (Word/Excel), đúng thứ người dùng nhận được.

Chạy: python tests/browser-real-pdf-test.py   (cần server xem trước ở :4173)
"""
import sys
from pathlib import Path

from playwright.sync_api import sync_playwright

REPO = Path(__file__).resolve().parent.parent
BASE = "http://localhost:4173"
BANG = REPO / "tests" / "real" / "bang-tong-hop.pdf"
CV = REPO / "tests" / "real" / "cv-5095-bgd.pdf"
OUT = REPO / "tests" / "real" / "out"
OUT.mkdir(parents=True, exist_ok=True)

failures = []


def check(desc, cond):
    print(("  ✓ " if cond else "  ✘ ") + desc)
    if not cond:
        failures.append(desc)


def wait_done(page, timeout=300000):
    """Chuyển xong = nút tải Word hiện (chỉ render khi đã có kết quả)."""
    btn = page.locator("button", has_text="Tải file Word")
    btn.wait_for(state="visible", timeout=timeout)
    return btn


def docx_full_text(doc):
    full = "\n".join(p.text for p in doc.paragraphs)
    for tb in doc.tables:
        for r in tb.rows:
            for c in r.cells:
                full += "\n" + c.text
    return full


def main() -> int:
    with sync_playwright() as pw:
        browser = pw.chromium.launch()
        page = browser.new_page(accept_downloads=True)
        errors = []
        page.on("pageerror", lambda e: errors.append(str(e)))
        page.on("console", lambda m: errors.append(m.text) if m.type == "error" else None)
        page.goto(BASE + "/", wait_until="networkidle")

        print("File 1 — bảng tổng hợp (PDF in từ hệ thống trường học, font CID)")
        page.set_input_files("#fileInput", str(BANG))
        wait_done(page, timeout=180000)

        with page.expect_download(timeout=60000) as dl_info:
            page.click("button:has-text('Tải file Excel')")
        dl = dl_info.value
        xlsx_path = OUT / "bang-tong-hop-app.xlsx"
        dl.save_as(str(xlsx_path))
        import openpyxl

        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
        all_cells = [
            str(c.value)
            for ws in wb.worksheets
            for row in ws.iter_rows()
            for c in row
            if c.value is not None
        ]
        joined = " ".join(all_cells)
        check("Excel tải xuống chứa 'Vừ Mí Bình'", "Vừ Mí Bình" in joined)
        check("Excel tải xuống không chứa rác 's5igs'", "s5igs" not in joined)

        with page.expect_download(timeout=60000) as dl_info:
            page.click("button:has-text('Tải file Word')")
        dl = dl_info.value
        docx1_path = OUT / "bang-tong-hop-app.docx"
        dl.save_as(str(docx1_path))
        from docx import Document

        full1 = docx_full_text(Document(docx1_path))
        check("Word có tên học sinh đọc đúng từ lớp chữ ('Vừ Mí Bình')", "Vừ Mí Bình" in full1)
        check("Word có 'Phẩm chất chủ yếu'", "Phẩm chất chủ yếu" in full1)
        check("Không còn ký tự rác kiểu OCR ('s5igs', 'wu')", "s5igs" not in full1 and "wu" not in full1)
        check(f"Dấu '|' lạ gần như không còn (đếm: {full1.count('|')})", full1.count("|") < 5)

        print("File 2 — CV 5095/BGDĐT (bản SCAN thật từ cổng ký số: đúng đường OCR)")
        page.set_input_files("#fileInput", str(CV))
        wait_done(page, timeout=300000)
        body = page.evaluate("() => document.body.innerText")
        check("Có cảnh báo các trang scan đã quét OCR", "quét OCR" in body)

        with page.expect_download(timeout=60000) as dl_info:
            page.click("button:has-text('Tải file Word')")
        dl = dl_info.value
        docx_path = OUT / "cv-5095-app.docx"
        dl.save_as(str(docx_path))
        doc = Document(docx_path)
        full = docx_full_text(doc)
        check("Word tải xuống từ bản scan có nội dung đọc được", "Nghị quyết" in full and len(full) > 800)
        check("OCR đọc được 'BỘ GIÁO DỤC' ở phần đầu", "BỘ GIÁO DỤC" in full)
        # Hồi quy lỗi 23/10/2026: khi tài liệu có trang AI chen trang OCR,
        # khâu ghép từng trang một khiến MỌI trang bị dựng "phần đầu thể
        # thức" giả (bảng 2x2) và NUỐT mất nội dung đầu trang 2-3. File này
        # (không khóa → toàn OCR) kiểm dạng thuần OCR; nhánh ghép hỗn hợp
        # được kiểm ở browser-ai-test. Ở đây bắt buộc: đúng 1 bảng (phần đầu
        # trang 1) và nội dung đầu trang 2 còn nguyên.
        check(f"Đúng 1 bảng trong Word (chỉ phần đầu trang 1) — thực tế {len(doc.tables)}", len(doc.tables) == 1)
        # (Probe tránh đúng chữ "tổ chức": Tesseract trên trình duyệt đọc
        # hụt dấu thành "t chức" — lỗi 1 ký tự của OCR, không phải mất nội
        # dung; hai mẩu dưới đây chứng minh mục 1 + hàng đầu bảng còn nguyên.)
        check("Nội dung mục 1 trang 2 còn nguyên ('cá nhân tham gia góp ý')", "cá nhân tham gia góp ý" in full)
        check("Hàng đầu bảng mục 1 còn ('Số lượng các xã')", "Số lượng các xã" in full)
        check("Tiêu đề mẫu báo cáo trang 2 còn ('BAO CÁO')", "BAO CÁO" in full)

        if errors:
            print("  (lỗi console/trang:", errors[:3], ")")
        browser.close()

    print()
    if failures:
        print("CHƯA ĐẠT:", failures)
        return 1
    print("TRÌNH DUYỆT + PDF THẬT: mọi kiểm tra ĐẠT.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
