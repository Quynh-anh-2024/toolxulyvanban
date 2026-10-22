"""Probe OCR trong TRÌNH DUYỆT THẬT: nạp CV 5095 (scan) vào app, lấy toàn bộ
text xem trước + file .docx tải xuống, đo độ phủ từ so với bản chép tay trang 1
(tests/real/truth-cv-p1.txt) và lưu text ra /tmp để soi tận nơi.

Chạy: python tests/ocr-browser-probe.py   (cần server xem trước ở :4173)
"""
import re
import sys
import unicodedata
from collections import Counter
from pathlib import Path

from playwright.sync_api import sync_playwright

REPO = Path(__file__).resolve().parent.parent
BASE = "http://localhost:4173"
CV = REPO / "tests" / "real" / "cv-5095-bgd.pdf"
TRUTH = REPO / "tests" / "real" / "truth-cv-p1.txt"

TOKEN_RE = re.compile(r"[0-9A-Za-zÀ-ỹĐđ]+")


def tokens(t: str) -> Counter:
    return Counter(TOKEN_RE.findall(unicodedata.normalize("NFC", t)))


def coverage(truth: Counter, out: Counter) -> float:
    total = sum(truth.values())
    hit = sum(min(v, out.get(k, 0)) for k, v in truth.items())
    return hit / total if total else 1.0


def main() -> int:
    truth = tokens(TRUTH.read_text(encoding="utf-8"))
    print(f"Truth trang 1: {sum(truth.values())} từ")
    with sync_playwright() as pw:
        browser = pw.chromium.launch()
        page = browser.new_page()
        errors = []
        page.on("pageerror", lambda e: errors.append(str(e)))
        page.goto(BASE + "/", wait_until="networkidle")
        page.set_input_files("input[type=file]", str(CV))
        page.wait_for_function(
            "() => document.body.innerText.includes('Trang OCR: 3')",
            timeout=300000,
        )
        page.wait_for_function(
            "() => document.body.innerText.includes('Kính gửi') || document.body.innerText.includes('Thực hiện')",
            timeout=120000,
        )
        page.wait_for_timeout(3000)
        text = page.evaluate("() => document.body.innerText")
        Path("/tmp/cv-browser-preview.txt").write_text(text, encoding="utf-8")
        cov = coverage(truth, tokens(text))
        print(f"XEM TRƯỚC trình duyệt: phủ trang 1 = {cov*100:.1f}%")

        with page.expect_download(timeout=60000) as dl_info:
            page.click("text=Tải File Word (.docx) Chuẩn NĐ 30")
        dl = dl_info.value
        docx_path = "/tmp/cv-browser.docx"
        dl.save_as(docx_path)
        from docx import Document

        doc = Document(docx_path)
        full = "\n".join(p.text for p in doc.paragraphs)
        for tb in doc.tables:
            for r in tb.rows:
                for c in r.cells:
                    full += "\n" + c.text
        Path("/tmp/cv-browser-docx.txt").write_text(full, encoding="utf-8")
        cov2 = coverage(truth, tokens(full))
        print(f"WORD tải xuống: phủ trang 1 = {cov2*100:.1f}%  (tổng {len(full)} ký tự)")
        if errors:
            print("  (lỗi trang:", errors[:3], ")")
        browser.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
